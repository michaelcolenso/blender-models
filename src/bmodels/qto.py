"""Quantity takeoff (QTO) from a built house model -> Excel workbook with live formulas.

Called by build.py after verification. Measured geometry is written as blue input cells
with a method + status on every row; everything derived (net siding, squares, casing LF,
totals) is an Excel formula so the sheet recalculates if an input is corrected.

Status vocabulary (same as the model): documented / provisional (measured from a drawing
at a verified scale) / assumed.
"""
import json, math
from pathlib import Path

import bmesh
from mathutils import Vector

IN = 0.0254
SF = 1 / (0.0254 ** 2 * 144)  # m^2 -> sf


def elev_name(nx, ny):
    """Plan normal (X right, Y toward front) -> elevation name."""
    if abs(nx) >= abs(ny):
        return "Right side" if nx > 0 else "Left side"
    return "Front" if ny > 0 else "Rear"


def simplify(ring, min_len=4.0):
    pts = [tuple(p) for p in ring]
    changed = True
    while changed and len(pts) > 4:
        changed = False
        for i in range(len(pts)):
            a, b, c = pts[i - 1], pts[i], pts[(i + 1) % len(pts)]
            l1 = math.dist(a, b); l2 = math.dist(b, c)
            cr = (b[0] - a[0]) * (c[1] - b[1]) - (b[1] - a[1]) * (c[0] - b[0])
            if l1 < min_len or abs(cr) < 1e-6 * max(l1 * l2, 1e-9) * 1e3:
                pts.pop(i); changed = True; break
    return pts


def _inside(pt, ring):
    x, y = pt; c = False
    for i in range(len(ring)):
        (x1, y1), (x2, y2) = ring[i - 1], ring[i]
        if (y1 > y) != (y2 > y) and x < (x2 - x1) * (y - y1) / (y2 - y1) + x1:
            c = not c
    return c


def walls_by_elevation(fp, H):
    """Exterior wall lengths per elevation from the footprint ring (right half + mirror).
    Outward side of each edge is found by a point-in-polygon test, so ring orientation
    does not matter."""
    rows, corners = [], 0
    for side, ring in (("R", fp), ("L", [(-x, y) for x, y in reversed(fp)])):
        r = simplify(ring)
        normals = []
        for i in range(len(r)):
            a, b = r[i - 1], r[i]
            dx, dy = b[0] - a[0], b[1] - a[1]; L = math.hypot(dx, dy) or 1
            nx, ny = dy / L, -dx / L
            mid = ((a[0] + b[0]) / 2 + nx * 3, (a[1] + b[1]) / 2 + ny * 3)
            if _inside(mid, r):
                nx, ny = -nx, -ny
            normals.append((nx, ny))
            if abs(a[0]) < 3 and abs(b[0]) < 3:
                continue  # party-wall line, not exterior
            rows.append({"unit": side, "elevation": elev_name(nx, ny), "length_in": round(L, 1)})
        for i in range(len(r)):
            a = r[i - 1]; b = r[i]; n1 = normals[i]; n2 = normals[(i + 1) % len(r)]
            if abs(b[0]) < 3:
                continue
            probe = (b[0] + (n1[0] + n2[0]) * 3, b[1] + (n1[1] + n2[1]) * 3)
            if abs(n1[0] * n2[1] - n1[1] * n2[0]) > 0.5 and not _inside(probe, r):
                corners += 1  # outward-convex corner
    return rows, corners


def roof_lines(obj):
    bm = bmesh.new(); bm.from_mesh(obj.data); bm.transform(obj.matrix_world)
    top = lambda f: f.normal.z > 0.3
    out = {"ridge": 0.0, "hip": 0.0, "valley": 0.0, "eave": 0.0, "rake": 0.0}
    pitch_area = {}
    for f in bm.faces:
        if top(f):
            p = round(math.tan(math.acos(min(1, f.normal.z))) * 12 * 2) / 2
            pa = pitch_area.setdefault(p, [0.0, 0.0])
            pa[0] += f.calc_area() * SF; pa[1] += f.calc_area() * f.normal.z * SF
    for e in bm.edges:
        tf = [f for f in e.link_faces if top(f)]
        L = e.calc_length() / IN / 12
        dz = abs(e.verts[0].co.z - e.verts[1].co.z) / IN
        if len(tf) == 2:
            if tf[0].normal.dot(tf[1].normal) > 0.9999:
                continue  # coplanar split, not a line
            if e.is_convex:
                out["ridge" if dz < 0.5 else "hip"] += L
            else:
                out["valley"] += L
        elif len(tf) == 1:
            other = [f for f in e.link_faces if f is not tf[0]]
            if other and abs(other[0].normal.z) < 0.2:  # top face meets fascia/rake board
                out["eave" if dz < 0.5 else "rake"] += L
    bm.free()
    return out, pitch_area


def gable_areas(objs):
    res = {}
    for o in objs:
        bm = bmesh.new(); bm.from_mesh(o.data); bm.transform(o.matrix_world)
        for f in bm.faces:
            if abs(f.normal.z) > 0.05:
                continue
            zs = [v.co.z / IN for v in f.verts]
            if max(zs) - min(zs) < 24:
                continue  # eave strip hidden behind soffit, not a gable
            ctr = f.calc_center_median() / IN
            if abs(ctr.x) < 6 and abs(f.normal.x) > 0.9:
                continue  # party-wall plane between the two attics — interior, not siding
            k = elev_name(f.normal.x, -f.normal.y)
            res[k] = res.get(k, 0.0) + f.calc_area() * SF
        bm.free()
    return res


NOMINAL = {  # AH500 cottage window schedule (documented)
    "A": "(3) 2-8 x 5-6 casement", "B": "(2) 2-8 x 5-6 casement",
    "C": "(1) 3-0 x 5-6 casement", "D": "(1) 2-6 x 5-6 casement",
}


def run(spec, built, out_dir, house_mod):
    el = next(e for e in spec["elements"] if e.get("type") == "house")
    tk = json.loads((Path(spec["_dir"]) / el["takeoff"]).read_text())
    H = el["plate_height_in"]
    wall_rows, corners = walls_by_elevation(tk["footprint"], H)
    roof = next(o for m, o in built if o.name == "roof")
    attics = [o for m, o in built if o.name.startswith("attic_")]
    lines, pitch_area = roof_lines(roof)
    gables = gable_areas(attics)
    ops = house_mod.mirror_openings(el["openings"])
    plan_total = sum(v[1] for v in pitch_area.values())

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.comments import Comment
    wb = Workbook()
    F = "Arial"; blue = Font(name=F, color="0000FF"); black = Font(name=F); bold = Font(name=F, bold=True)
    hdr_fill = PatternFill("solid", fgColor="DDE3EA"); yellow = PatternFill("solid", fgColor="FFFF00")

    def header(ws, cols, widths):
        ws.append(cols)
        for c, w in zip(ws[1], widths):
            c.font = bold; c.fill = hdr_fill; c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.column_dimensions[c.column_letter].width = w
        ws.freeze_panes = "A2"

    # ---------- Openings ----------
    wo = wb.active; wo.title = "Openings"
    header(wo, ["ID", "Unit", "Elevation", "Mark", "Kind", "Nominal (AH500 schedule)", "Width (in)",
                "Sill (in)", "Head (in)", "Height (in)", "Area (sf)", "Casing (LF)", "Status", "Source / basis"],
           [16, 6, 11, 6, 8, 24, 10, 8, 8, 10, 10, 11, 12, 48])
    for o in ops:
        unit = "R" if (o["x0"] + o["x1"]) > 0 else "L"
        if o["axis"] == "x":
            el_ = "Rear" if o.get("face", 1) < 0 else "Front"; w = o["x1"] - o["x0"]
        else:
            el_ = "Right side" if o.get("face", 1) > 0 else "Left side"; w = o["y1"] - o["y0"]
        r = wo.max_row + 1
        mark = o.get("mark", "")
        wo.append([o["id"], unit, el_, mark, o["kind"], NOMINAL.get(mark, ""), round(w, 1), o["sill"], o["head"],
                   f"=I{r}-H{r}", f"=G{r}*J{r}/144",
                   f"=2*(G{r}+J{r})/12" if o["kind"] == "window" else f"=(G{r}+2*J{r})/12",
                   o.get("status", "provisional"),
                   (o.get("basis") or f"width: AH201 plan extraction @2px/in; heights: {o.get('source', 'AH500')} scaled"),
                   ])
        for col in "GHI":
            wo[f"{col}{r}"].font = blue
    n_op = wo.max_row
    wo.append([]); wo.append(["Totals"]); tr = wo.max_row
    wo[f"A{tr}"].font = bold
    wo[f"K{tr}"] = f"=SUM(K2:K{n_op})"; wo[f"L{tr}"] = f"=SUM(L2:L{n_op})"
    wo.append(["Widths are the rough gap measured in the plan (frame to frame), not catalog unit sizes."])

    # ---------- Walls ----------
    ww = wb.create_sheet("Walls")
    header(ww, ["Elevation", "Unit", "Wall length (in)", "Wall height (in)", "Gross wall (sf)", "Status", "Method"],
           [12, 6, 14, 14, 14, 12, 60])
    for row in wall_rows:
        r = ww.max_row + 1
        ww.append([row["elevation"], row["unit"], row["length_in"], H, f"=C{r}*D{r}/144", "provisional",
                   "footprint edge from AH201 extraction (±1 in) × FF-to-truss-bearing 109 1/8\" (documented)"])
        ww[f"C{r}"].font = blue; ww[f"D{r}"].font = blue
    n_w = ww.max_row
    # summary by elevation
    ww.append([]); ww.append(["By elevation", "", "Gross wall (sf)", "Gable panel FP (sf)", "Openings (sf)",
                              "Net siding + panel (sf)", "Method"])
    for c in ww[ww.max_row]:
        c.font = bold
    elevs = ["Front", "Rear", "Right side", "Left side"]
    first = ww.max_row + 1
    for e in elevs:
        r = ww.max_row + 1
        ww.append([e, "", f'=SUMIFS(E2:E{n_w},A2:A{n_w},A{r})', round(gables.get(e, 0.0), 1),
                   f'=SUMIFS(Openings!K2:K{n_op},Openings!C2:C{n_op},A{r})', f"=C{r}+D{r}-E{r}",
                   "gable: attic faces at wall plane taller than 24 in (model); openings deducted at rough size"])
        ww[f"D{r}"].font = blue
    last = ww.max_row
    ww.append(["Total", "", f"=SUM(C{first}:C{last})", f"=SUM(D{first}:D{last})", f"=SUM(E{first}:E{last})",
               f"=SUM(F{first}:F{last})"])
    tot_w = ww.max_row
    for c in ww[tot_w]:
        c.font = bold
    ww.append([]); ww.append(["Outside corners (count)", "", corners, "", "", "", "convex footprint corners after removing <4 in jogs"])
    cr_ = ww.max_row; ww[f"C{cr_}"].font = blue
    ww.append(["Corner boards (LF)", "", f"=C{cr_}*{H}/12", "", "", "", "corners × wall height; 5/4x4 per AH500 notes"])
    cb_ = ww.max_row

    # ---------- Roof ----------
    wr = wb.create_sheet("Roof")
    header(wr, ["Item", "Pitch", "Quantity", "Unit", "Status", "Method / source"], [26, 8, 12, 8, 12, 62])
    wr.append(["Waste factor (input)", "", 0.10, "", "assumed", "EDIT: typical allowance for hips/valleys; not from documents"])
    wr["C2"].font = blue; wr["C2"].fill = yellow; wr["C2"].number_format = "0%"
    wr["C2"].comment = Comment("Assumption — change to your estimating standard.", "QTO")
    pa_rows = []
    for p in sorted(pitch_area):
        if pitch_area[p][0] < 1:
            continue  # sliver faces from boolean seams
        r = wr.max_row + 1
        wr.append([f"Sloped area {p:g}/12", f"{p:g}/12", round(pitch_area[p][0], 1), "sf", "provisional",
                   "sum of roof top faces in model by pitch (pitches documented AH202)"])
        wr[f"C{r}"].font = blue; pa_rows.append(r)
    r = wr.max_row + 1
    wr.append(["Sloped area total", "", f"=SUM(C{pa_rows[0]}:C{pa_rows[-1]})", "sf", "provisional", ""]); tot_r = r
    wr.append(["Squares (with waste)", "", f"=C{tot_r}*(1+C2)/100", "sq", "provisional", "sloped area × (1 + waste) ÷ 100"])
    sq_r = wr.max_row
    r = wr.max_row + 1
    wr.append(["Plan (horizontal) area", "", round(plan_total, 1), "sf", "provisional",
               "model eave outline; AH202 venting calc states 3,935 sf of roof area (documented)"])
    wr[f"C{r}"].font = blue; plan_r = r
    wr.append(["  difference vs AH202 3,935 sf", "", f"=C{plan_r}-3935", "sf", "", "check"])
    for k, lbl in [("ridge", "Ridge"), ("hip", "Hip"), ("valley", "Valley"), ("eave", "Eave (gutter/fascia)"),
                   ("rake", "Rake")]:
        r = wr.max_row + 1
        wr.append([lbl, "", round(lines[k], 1), "LF", "provisional", "roof mesh edges classified by convexity/slope"])
        wr[f"C{r}"].font = blue
    ridge_r = sq_r + 3

    # ---------- Summary ----------
    s = wb.create_sheet("Summary", 0)
    header(s, ["Quantity", "Value", "Unit", "Where"], [34, 14, 8, 40])
    s.append([spec.get("project", ""), "", "", "Generated by src/bmodels/qto.py from the verified model"])
    s["A2"].font = bold
    rows = [
        ("Windows (count)", f'=COUNTIFS(Openings!E2:E{n_op},"window")', "ea", "Openings"),
        ("Exterior doors (count)", f'=COUNTIFS(Openings!E2:E{n_op},"door")', "ea", "Openings"),
        ("Overhead garage doors (count)", f'=COUNTIFS(Openings!E2:E{n_op},"garage")', "ea", "Openings"),
        ("Net siding + gable panel", f"=Walls!F{tot_w}", "sf", "Walls — by elevation"),
        ("  of which gable panel (FP)", f"=Walls!D{tot_w}", "sf", "Walls"),
        ("Opening casing trim", f"=Openings!L{tr}", "LF", "Openings"),
        ("Corner boards", f"=Walls!C{cb_}", "LF", "Walls"),
        ("Roof sloped area", f"=Roof!C{tot_r}", "sf", "Roof"),
        ("Roofing squares incl. waste", f"=Roof!C{sq_r}", "sq", "Roof (waste is an input)"),
        ("Roof plan area (AH202 says 3,935 sf)", f"=Roof!C{plan_r}", "sf", "Roof"),
    ]
    for q in rows:
        s.append(list(q))
    for row in s.iter_rows(min_row=3, max_row=s.max_row):
        row[1].number_format = "#,##0.0"
    s.append([]); s.append(["Legend"]); s[f"A{s.max_row}"].font = bold
    for t in ["Blue = measured/input value (edit to correct); black = formula.",
              "Yellow = assumption you should set (roof waste).",
              "Status: documented = shown on a sheet; provisional = measured from a drawing at a verified scale;",
              "  assumed = not in documents (see basis). Openings are rough gap sizes, not catalog sizes.",
              "Not included: pergolas/porches, bay mini-roofs, gutters/downspouts, trim bands, interior."]:
        s.append([t])

    # ---------- Elements (tracker keys) ----------
    we = wb.create_sheet("Elements")
    header(we, ["Element ID", "Trade", "Status", "Source", "Basis"], [28, 12, 12, 10, 60])
    for m, o in built:
        k = o.name
        if m.get("trade") == "openings":
            if not k.endswith(("_glass", "_panel")):
                continue  # one row per opening, not per trim piece
            k = k.rsplit("_", 1)[0]
        we.append([k, m.get("trade", ""), m.get("status", ""), m.get("source", ""), m.get("basis") or ""])

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.font is None or c.font.name != F:
                    c.font = Font(name=F, bold=c.font.bold if c.font else False,
                                  color=c.font.color if c.font else None)
    out = Path(out_dir) / f"{Path(spec['_path']).stem}_takeoff.xlsx"
    wb.save(out)
    return out, {"plan_area_sf": round(plan_total, 1), "lines": lines, "gables": gables, "corners": corners}
