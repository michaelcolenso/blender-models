"""Slab-cut, excavation and pipe quantities for the Building B under-slab re-pipe.

Reads meydenbauer-b.spec.json (from prepare.py), computes the trench footprint as the union of
the drawn routes widened to the trench width, then sizes excavation from the invert profile.
Writes takeoff.json (consumed by build.py) and Meydenbauer-B-underslab-takeoff.xlsx.

Usage: python projects/meydenbauer-b/products/underslab/takeoff.py [--out build/meydenbauer-b]
"""
import argparse, json, math
from collections import defaultdict
from pathlib import Path

import numpy as np
import shapely
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from shapely.geometry import LineString, Point, box
from shapely.ops import unary_union

HERE = Path(__file__).parent
from bmodels.paths import ROOT as REPO  # noqa: E402
SPEC = json.loads((HERE / "meydenbauer-b.spec.json").read_text())

ALLOWANCE_7_SF = 700.0              # Bid Form 00 41 13 (Addendum 1) Allowance No. 7
SLOPE = 0.25 / 12                   # ft/ft, "1/4 SLOPE" on P1.02 and General Note
BEDDING_FT = 4 / 12                 # 2/P5.0: 4" bedding below pipe
OUTSIDE_FT = 1.5                    # SOR: replacement starts 18" outside the footprint
# OD by nominal size. 4" and 6" from SOR Figs 3-4 (Viterra ABS); 2" and 3" are the standard
# IPS ODs for Sch 40 DWV (assumed — not stated in the bid documents).
OD_IN = {2: 2.375, 3: 3.5, 4: 4.5, 6: 6.625}
OD_STATUS = {2: "assumed", 3: "assumed", 4: "confirmed", 6: "confirmed"}

SCENARIOS = [
    # id, label, trench width ft, pit side ft (0 = none), widen-to ft where trench depth >= 4 ft
    ("S1", "Detail 10/P5.1 typical: 2'-0\" trench, no extra pits", 2.0, 0.0, None),
    ("S2", "S1 + 3'x3' working pit at every fixture/stack connection", 2.0, 3.0, None),
    ("S3", "S2 + 4'-0\" wide where trench is 4'-0\" deep or more", 2.0, 3.0, 4.0),
]
TIE_DEPTHS = [2.5, 3.5, 4.5]        # ft below finished floor at the POC; SOR recorded 2-1/2 to 4-1/2 ft
SLAB_FT = 4 / 12                    # existing slab thickness: not documented (assumed 4")
SWELL = 1.25                        # bank-to-loose (assumed)


def unit_geometry(u):
    N = u["nodes"]
    segs = []
    for e in u["edges"]:
        a, b = N[e["a"]], N[e["b"]]
        segs.append({"a": a["xy_ft"], "b": b["xy_ft"], "fa": a["flow_ft"], "fb": b["flow_ft"],
                     "nominal": e["nominal_in"], "L": e["length_ft"]})
    ends = [{"xy": N[e["node"]]["xy_ft"], "kind": e["kind"], "flow": N[e["node"]]["flow_ft"]}
            for e in u["ends"] if e["kind"] != "TEE"]
    tie = u["tie_in"]["xy_ft"]
    return segs, ends, tie


def outside_zone(segs, tie):
    """Run within OUTSIDE_FT of flow from the POC: outside the footprint (landscape, not slab)."""
    parts = []
    for s in segs:
        lo, hi = sorted((s["fa"], s["fb"]))
        if lo >= OUTSIDE_FT:
            continue
        a, b = (s["a"], s["b"]) if s["fa"] <= s["fb"] else (s["b"], s["a"])
        t = min(1.0, (OUTSIDE_FT - lo) / max(hi - lo, 1e-9))
        parts.append(LineString([a, (a[0] + (b[0] - a[0]) * t, a[1] + (b[1] - a[1]) * t)]))
    return unary_union(parts).buffer(1.5, cap_style="square")    # generous mask, cut back below


def depth_field(segs, tie_depth):
    """Dense samples (x, y, invert depth) along every run."""
    pts = []
    for s in segs:
        n = max(2, int(s["L"] / 0.1) + 1)
        for t in np.linspace(0, 1, n):
            x = s["a"][0] + (s["b"][0] - s["a"][0]) * t
            y = s["a"][1] + (s["b"][1] - s["a"][1]) * t
            f = s["fa"] + (s["fb"] - s["fa"]) * t
            pts.append((x, y, tie_depth - SLOPE * f))
    return np.array(pts)


def footprint(segs, ends, width, pit, widen=None, deep_mask=None):
    lines = [LineString([s["a"], s["b"]]) for s in segs]
    geom = unary_union([l.buffer(width / 2, cap_style="square", join_style="mitre") for l in lines])
    if pit:
        geom = unary_union([geom] + [box(e["xy"][0] - pit / 2, e["xy"][1] - pit / 2,
                                         e["xy"][0] + pit / 2, e["xy"][1] + pit / 2) for e in ends])
    if widen and deep_mask is not None and not deep_mask.is_empty:
        geom = unary_union([geom, deep_mask.buffer(widen / 2, cap_style="flat", join_style="mitre")])
    return geom


def deep_runs(segs, tie_depth, threshold=4.0):
    """Centreline pieces whose trench depth (invert + bedding) is >= threshold."""
    parts = []
    for s in segs:
        da = tie_depth - SLOPE * s["fa"] + BEDDING_FT
        db = tie_depth - SLOPE * s["fb"] + BEDDING_FT
        if max(da, db) < threshold:
            continue
        if min(da, db) >= threshold:
            parts.append(LineString([s["a"], s["b"]])); continue
        t = (threshold - da) / (db - da)
        m = (s["a"][0] + (s["b"][0] - s["a"][0]) * t, s["a"][1] + (s["b"][1] - s["a"][1]) * t)
        parts.append(LineString([s["a"], m]) if da >= threshold else LineString([m, s["b"]]))
    return unary_union(parts) if parts else LineString()


def excavation_cy(geom, field, cell=0.25):
    """Bank volume below the slab: each footprint cell takes the depth of its nearest run."""
    x0, y0, x1, y1 = geom.bounds
    xs = np.arange(x0 + cell / 2, x1, cell); ys = np.arange(y0 + cell / 2, y1, cell)
    X, Y = np.meshgrid(xs, ys); X = X.ravel(); Y = Y.ravel()
    inside = shapely.contains_xy(geom, X, Y)
    X, Y = X[inside], Y[inside]
    if not len(X):
        return 0.0, 0.0
    d2 = (X[:, None] - field[None, :, 0]) ** 2 + (Y[:, None] - field[None, :, 1]) ** 2
    depth = field[d2.argmin(axis=1), 2] + BEDDING_FT - SLAB_FT
    vol_cf = float(np.clip(depth, 0, None).sum() * cell * cell)
    return vol_cf / 27.0, float(depth.max() + SLAB_FT)


def main(out):
    out.mkdir(parents=True, exist_ok=True)
    results = {"inputs": {"allowance_7_sf": ALLOWANCE_7_SF, "slope_ft_per_ft": SLOPE,
                          "bedding_ft": BEDDING_FT, "slab_ft": SLAB_FT, "swell": SWELL,
                          "outside_ft": OUTSIDE_FT, "tie_depths_ft": TIE_DEPTHS,
                          "scenarios": [list(s) for s in SCENARIOS]},
               "units": []}
    for u in SPEC["units"]:
        segs, ends, tie = unit_geometry(u)
        outside = outside_zone(segs, tie)
        by_size = defaultdict(float)
        for s in segs:
            by_size[s["nominal"]] += s["L"]
        kinds = defaultdict(int)
        for e in ends:
            kinds[e["kind"]] += 1
        max_flow = max(n["flow_ft"] for n in u["nodes"].values())
        ur = {"unit": u["unit"], "pipe_lf": sum(by_size.values()),
              "pipe_lf_by_size": {str(k): v for k, v in sorted(by_size.items())},
              "ends": dict(kinds), "connections": len(ends), "max_flow_ft": max_flow,
              "fall_ft": SLOPE * max_flow, "cases": {}}
        for d0 in TIE_DEPTHS:
            field = depth_field(segs, d0)
            deep = deep_runs(segs, d0)
            for sid, label, w, pit, widen in SCENARIOS:
                g = footprint(segs, ends, w, pit, widen, deep)
                slab = g.difference(outside)
                cy, dmax = excavation_cy(g, field)
                ur["cases"][f"{sid}@{d0}"] = {
                    "slab_cut_sf": slab.area, "outside_sf": g.area - slab.area,
                    "excavation_bank_cy": cy, "spoil_loose_cy": cy * SWELL,
                    "max_trench_depth_ft": dmax, "deep_lf": deep.length,
                    "min_invert_depth_ft": d0 - SLOPE * max_flow,
                }
                if sid == "S2" and d0 == 3.5:
                    ur["footprint_S2"] = [list(p.exterior.coords) for p in getattr(slab, "geoms", [slab])]
                    ur["outside_S2"] = [list(p.exterior.coords) for p in
                                        getattr(g.intersection(outside), "geoms", [g.intersection(outside)])
                                        if p.geom_type == "Polygon"]
        results["units"].append(ur)

    tot = lambda key, case: sum(u["cases"][case][key] for u in results["units"])
    results["totals"] = {case: {"slab_cut_sf": tot("slab_cut_sf", case),
                                "excavation_bank_cy": tot("excavation_bank_cy", case),
                                "spoil_loose_cy": tot("spoil_loose_cy", case),
                                "deep_lf": tot("deep_lf", case)}
                         for case in results["units"][0]["cases"]}
    results["totals"]["pipe_lf"] = sum(u["pipe_lf"] for u in results["units"])
    (out / "takeoff.json").write_text(json.dumps(results, indent=1))
    write_xlsx(results, out / "Meydenbauer-B-underslab-takeoff.xlsx")

    print(f'pipe LF total {results["totals"]["pipe_lf"]:.1f}')
    for case in results["units"][0]["cases"]:
        t = results["totals"][case]
        print(f'{case:8s} slab cut {t["slab_cut_sf"]:6.0f} SF  ({t["slab_cut_sf"] - ALLOWANCE_7_SF:+5.0f} vs 700)'
              f'  excavation {t["excavation_bank_cy"]:5.1f} CY bank  deep(>=4ft) {t["deep_lf"]:5.1f} LF')
    for u in results["units"]:
        c = u["cases"]["S2@3.5"]
        print(f'{u["unit"]}: {u["pipe_lf"]:5.1f} LF {u["pipe_lf_by_size"]}  S2@3.5 slab {c["slab_cut_sf"]:5.0f} SF'
              f'  min invert {c["min_invert_depth_ft"]:.2f} ft  ends {u["ends"]}')


BLUE = Font(color="1F4E9E")
HEAD = Font(bold=True)
FILL = PatternFill("solid", fgColor="FFF2CC")


def write_xlsx(r, path):
    wb = Workbook()
    ws = wb.active; ws.title = "Summary"
    ws.append(["101 Meydenbauer Bldg B — under-slab waste re-pipe: slab-cut check against Allowance 7"]); ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Estimating aid derived from bid set P1.02 linework. Not a contract document. Blue = measured/assumed input; black = formula."])
    ws.append([])
    ws.append(["Allowance No. 7 (Bid Form, Addendum 1)", ALLOWANCE_7_SF, "SF", "concrete cut, excavate, backfill, vapor barrier, crystalline additive, patch"])
    ws["B4"].font = BLUE
    ws.append([])
    hdr = ["Case", "Trench / pit basis", "POC invert depth (ft)", "Slab cut SF (measured)", "vs Allowance 7 (SF)",
           "Excavation bank CY (measured)", "Spoil loose CY", "Trench >= 4 ft deep (LF)"]
    ws.append(hdr)
    for c in ws[ws.max_row]:
        c.font = HEAD
    labels = {s[0]: s[1] for s in SCENARIOS}
    for case, t in r["totals"].items():
        if case == "pipe_lf":
            continue
        sid, d0 = case.split("@")
        row = ws.max_row + 1
        ws.append([case, labels[sid], float(d0), round(t["slab_cut_sf"], 1), f"=D{row}-$B$4",
                   round(t["excavation_bank_cy"], 1), f"=F{row}*Inputs!$B$6", round(t["deep_lf"], 1)])
        for col in "CDFH":
            ws[f"{col}{row}"].font = BLUE
        if case == "S2@3.5":
            for c in ws[row]:
                c.fill = FILL
    ws.append([])
    ws.append(["Highlighted row = base case used for the 3D model (S2, POC invert 3.5 ft)."])
    for col, w in zip("ABCDEFGH", [10, 58, 12, 14, 14, 16, 12, 14]):
        ws.column_dimensions[col].width = w

    wu = wb.create_sheet("By unit")
    wu.append(["Unit", "Pipe LF total", "2\" LF", "3\" LF", "4\" LF", "Connections (fixture/stack ends)",
               "Longest flow path (ft)", "Fall at 1/4\"/ft (ft)", "S1@3.5 slab SF", "S2@3.5 slab SF", "S3@3.5 slab SF",
               "S2@3.5 excavation CY", "Min invert depth @3.5 (ft)", "Ends"])
    for c in wu[1]:
        c.font = HEAD; c.alignment = Alignment(wrap_text=True)
    for i, u in enumerate(r["units"], start=2):
        ls = u["pipe_lf_by_size"]
        wu.append([u["unit"], f"=C{i}+D{i}+E{i}", round(ls.get("2", 0), 1), round(ls.get("3", 0), 1), round(ls.get("4", 0), 1),
                   u["connections"], round(u["max_flow_ft"], 1), f"=G{i}*Inputs!$B$3",
                   round(u["cases"]["S1@3.5"]["slab_cut_sf"], 1), round(u["cases"]["S2@3.5"]["slab_cut_sf"], 1),
                   round(u["cases"]["S3@3.5"]["slab_cut_sf"], 1), round(u["cases"]["S2@3.5"]["excavation_bank_cy"], 1),
                   round(u["cases"]["S2@3.5"]["min_invert_depth_ft"], 2), json.dumps(u["ends"])])
        for col in "CDEFGIJKLM":
            wu[f"{col}{i}"].font = BLUE
    n = len(r["units"]) + 1
    wu.append(["Total"] + [f"=SUM({c}2:{c}{n})" for c in "BCDEF"] + ["", ""] + [f"=SUM({c}2:{c}{n})" for c in "IJKL"])
    for c in wu[wu.max_row]:
        c.font = HEAD
    for col in "ABCDEFGHIJKLMN":
        wu.column_dimensions[col].width = 13
    wu.column_dimensions["N"].width = 60

    wi = wb.create_sheet("Inputs")
    rows = [
        ("Input", "Value", "Unit", "Status", "Source / basis"),
        ("Allowance 7", ALLOWANCE_7_SF, "SF", "confirmed", "Bid Form 00 41 13 Addendum 1, Allowance No. 7"),
        ("Slope", SLOPE, "ft/ft", "confirmed", "P1.02 \"1/4 SLOPE\"; Construction Note 3"),
        ("Bedding below pipe", BEDDING_FT, "ft", "confirmed", "2/P5.0 bedding detail, 4\""),
        ("Trench width (typ.)", 2.0, "ft", "confirmed", "10/P5.1 Trench Infill, 2'-0\" TYP. (1'-3\" MIN.)"),
        ("Swell factor", SWELL, "-", "assumed", "bank to loose, typical for sandy/silty fill"),
        ("Existing slab thickness", SLAB_FT, "ft", "assumed", "not documented; pour-back is thicker of existing or 5\" (10/P5.1)"),
        ("Working pit at each connection", 3.0, "ft square", "assumed", "fittings, drain bodies, carriers, stack-base cleanouts"),
        ("Deep-trench width", 4.0, "ft", "assumed", "protective system/benching where trench >= 4 ft (verify WA DOSH, WAC 296-155 Part N)"),
        ("Outside-footprint length at POC", OUTSIDE_FT, "ft", "confirmed", "SOR: replace from 18\" outside the footprint inward"),
        ("POC invert depth cases", ", ".join(str(d) for d in TIE_DEPTHS), "ft below FF", "assumed", "SOR: average recorded depths 2-1/2 to 4-1/2 ft; actual inverts not documented (MATCH EX. I.E.)"),
        ("Plan scale", SPEC["calibration"]["pt_per_ft_assumed"], "pt/ft", "provisional",
         f'measured {SPEC["calibration"]["measured_x_pt_per_ft"]:.3f} (grid A-M) / {SPEC["calibration"]["measured_y_pt_per_ft"]:.3f} (grid 1-6)'),
    ]
    for row in rows:
        wi.append(list(row))
    for c in wi[1]:
        c.font = HEAD
    for i in range(2, len(rows) + 1):
        wi[f"B{i}"].font = BLUE
    for col, w in zip("ABCDE", [32, 14, 12, 12, 90]):
        wi.column_dimensions[col].width = w

    we = wb.create_sheet("Runs")
    we.append(["ID", "Unit", "From (x,y ft)", "To (x,y ft)", "Length ft", "Nominal in", "OD in", "OD status", "Fixtures served", "Status", "Source"])
    for c in we[1]:
        c.font = HEAD
    for u in SPEC["units"]:
        N = u["nodes"]
        for i, e in enumerate(u["edges"]):
            a = [round(v, 2) for v in N[e["a"]]["xy_ft"]]; b = [round(v, 2) for v in N[e["b"]]["xy_ft"]]
            we.append([f'{u["unit"]}-R{i:03d}', u["unit"], str(a), str(b), round(e["length_ft"], 2), e["nominal_in"],
                       OD_IN[e["nominal_in"]], OD_STATUS[e["nominal_in"]], e["fixtures_served"], "provisional", "P1.02 linework"])
    for col, w in zip("ABCDEFGHIJK", [14, 8, 18, 18, 10, 10, 8, 10, 10, 11, 16]):
        we.column_dimensions[col].width = w
    wb.save(path)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(REPO / "build" / "meydenbauer-b"))
    main(Path(ap.parse_args().out))
